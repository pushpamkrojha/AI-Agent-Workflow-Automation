from pathlib import Path
import openpyxl
from .models import Workflow

class WorkflowRegistry:
    """Loads workflow definitions from the Excel source at runtime."""
    def __init__(self, excel_path: str):
        self.excel_path = Path(excel_path)
        self.workflows: dict[str, Workflow] = {}
        self.load()

    def load(self):
        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        ws = wb["Workflows"]
        headers = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            wf = Workflow(
                workflow_id=data["Workflow_ID"], name=data["Workflow_Name"],
                trigger=data["Trigger"], inputs=data["Inputs"], steps=data["Steps"],
                decision_logic=data["Decision_Logic"], tools_required=data["Tools_Required"],
                expected_output=data["Expected_Output"],
            )
            self.workflows[wf.workflow_id] = wf

    def all(self) -> list[Workflow]:
        return list(self.workflows.values())

    def get(self, workflow_id: str) -> Workflow:
        return self.workflows[workflow_id]
